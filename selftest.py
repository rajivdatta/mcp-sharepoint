"""Offline self-test: exercises extraction, chunking, indexing and hybrid search
against synthetic files, with no SharePoint or network access.

Run:  .venv\\Scripts\\python.exe selftest.py
"""
from __future__ import annotations

import io
import json
import sys

import extract
import index as idx

OK, FAIL = "  ok  ", " FAIL "
results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(f"[{OK if cond else FAIL}] {name}" + (f"  - {detail}" if detail else ""))


# --------------------------------------------------------------------------- #
# 1. structured profiling
# --------------------------------------------------------------------------- #
csv_bytes = (
    "LotId,Community,CloseDate,Price,Status\n"
    "L-001,Preserve West,2026-03-14,742000,Closed\n"
    "L-002,Preserve West,2026-04-02,688500,Closed\n"
    "L-003,Cornell North,,915000,Firm\n"
    "L-004,Cornell North,2026-05-19,801250,Closed\n"
).encode()

prof = extract.profile_structured(csv_bytes, "csv", "closings_2026.csv")
t = prof["tables"][0]
check("csv profiled", t["rows"] == 4 and t["columns"] == 5, f"{t['rows']}x{t['columns']}")
lot = next(c for c in t["schema"] if c["name"] == "LotId")
check("candidate key detected", lot.get("candidate_key") is True)
close = next(c for c in t["schema"] if c["name"] == "CloseDate")
check("null % computed", close["nulls_pct"] == 25.0, f"{close['nulls_pct']}%")
price = next(c for c in t["schema"] if c["name"] == "Price")
check("numeric min/max", price.get("min") == 688500 and price.get("max") == 915000)

card = extract.schema_card(prof, "closings_2026.csv", "Sales/closings_2026.csv")
check("schema card mentions columns", "Community" in card and "CloseDate" in card)

# pipe-delimited (a common vendor-feed shape) must sniff correctly
pipe = b"a|b|c\n1|2|3\n4|5|6\n"
p2 = extract.profile_structured(pipe, "csv", "pipe.csv")
check("delimiter sniffed", p2["delimiter"] == "|" and p2["tables"][0]["columns"] == 3,
      p2["delimiter"])

# xlsx, multi-sheet
try:
    import pandas as pd
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"Region": ["GTA", "Calgary"], "Units": [12, 7]}).to_excel(
            w, sheet_name="Summary", index=False)
        pd.DataFrame({"Note": ["draft"]}).to_excel(w, sheet_name="Notes", index=False)
    p3 = extract.profile_structured(buf.getvalue(), "xlsx", "regions.xlsx")
    check("xlsx multi-sheet", len(p3["tables"]) == 2 and p3["sheet_count"] == 2,
          str([t["name"] for t in p3["tables"]]))
except Exception as exc:
    check("xlsx multi-sheet", False, str(exc))

# --------------------------------------------------------------------------- #
# 2. unstructured extraction
# --------------------------------------------------------------------------- #
try:
    import pymupdf
    doc = pymupdf.open()
    for i, body in enumerate([
            "Quarterly Land Review\nThe possession schedule for Cornell North "
            "slipped by six weeks because of servicing delays.",
            "Appendix A\nContact the vendor portal for the servicing agreement."]):
        page = doc.new_page()
        page.insert_text((72, 100), body, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    d = extract.extract_text(pdf_bytes, "pdf", "land_review.pdf")
    check("pdf pages extracted", len(d["pages"]) == 2, f"{len(d['pages'])} pages")
    check("pdf text extracted", "possession schedule" in d["pages"][0]["text"])
except Exception as exc:
    check("pdf extraction", False, str(exc))

try:
    import docx
    dd = docx.Document()
    dd.add_heading("Data Governance Standard", level=1)
    dd.add_paragraph("All datasets must carry an owner and a retention class.")
    tb = dd.add_table(rows=2, cols=2)
    tb.cell(0, 0).text = "Field"
    tb.cell(0, 1).text = "Owner"
    tb.cell(1, 0).text = "LotId"
    tb.cell(1, 1).text = "Land Analytics"
    b = io.BytesIO()
    dd.save(b)
    d2 = extract.extract_text(b.getvalue(), "docx", "standard.docx")
    txt = d2["pages"][0]["text"]
    check("docx headings marked", "## Data Governance Standard" in txt)
    check("docx table flattened", "LotId | Land Analytics" in txt)
except Exception as exc:
    check("docx extraction", False, str(exc))

try:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Retention policy"
    b = io.BytesIO()
    prs.save(b)
    d3 = extract.extract_text(b.getvalue(), "pptx", "deck.pptx")
    check("pptx slides extracted", "Retention policy" in d3["pages"][0]["text"])
except Exception as exc:
    check("pptx extraction", False, str(exc))

# --------------------------------------------------------------------------- #
# 3. chunking
# --------------------------------------------------------------------------- #
pages = [{"page": 1, "text": ("Alpha paragraph. " * 60 + "\n\n" +
                              "Beta paragraph. " * 60)},
         {"page": 2, "text": "Gamma paragraph about possession dates. " * 20}]
chunks = extract.chunk_pages(pages, target_chars=600, overlap=80)
check("chunks produced", len(chunks) >= 3, f"{len(chunks)} chunks")
check("chunks never straddle pages",
      all(c["page"] in (1, 2) for c in chunks) and
      {c["page"] for c in chunks} == {1, 2})
check("chunk size bounded", max(len(c["text"]) for c in chunks) <= 900,
      f"max {max(len(c['text']) for c in chunks)}")

# --------------------------------------------------------------------------- #
# 4. risk sweep
# --------------------------------------------------------------------------- #
risky = ("Login to the vendor portal\nuserid: jane.doe\npassword: Sup3rSecret!99\n"
         "AKIAIOSFODNN7EXAMPLE was the old key.")
hits = extract.scan_risks(risky)
kinds = {h["type"] for h in hits}
check("password detected", "password" in kinds)
check("aws key detected", "aws_key" in kinds)
check("secret redacted in context",
      all("Sup3rSecret" not in h["context"] and "AKIAIOSFODNN7EXAMPLE" not in h["context"]
          for h in hits),
      str(hits)[:140])
check("redact() is exhaustive",
      "Sup3rSecret" not in extract.redact(risky)
      and "AKIAIOSFODNN7EXAMPLE" not in extract.redact(risky))

# --------------------------------------------------------------------------- #
# 5. index + hybrid search (lexical path; no embedder)
# --------------------------------------------------------------------------- #
idx.DB_PATH = idx.DB_PATH.with_name("selftest.db")
idx.reset()

corpus = [
    ("land_review.pdf", "unstructured", "Land/land_review.pdf", [
        {"page": 1, "text": "The possession schedule for Cornell North slipped "
                            "by six weeks because of servicing delays."},
        {"page": 2, "text": "Appendix A lists the servicing agreement contacts."}]),
    ("closings_2026.csv", "structured", "Sales/closings_2026.csv", [
        {"page": None, "text": card}]),
    ("standard.docx", "unstructured", "Governance/standard.docx", [
        {"page": 1, "text": "All datasets must carry an owner and a retention class."}]),
]
for i, (name, kind, path, chs) in enumerate(corpus):
    item = {"item_id": f"item-{i}", "name": name, "path": path, "size": 1000,
            "ext": name.rsplit(".", 1)[-1], "modified": "2026-09-01T00:00:00Z",
            "created": "2026-01-01T00:00:00Z", "etag": f"tag-{i}",
            "web_url": f"https://example/{path}", "hash": None,
            "created_by": "Jane", "modified_by": "Jane"}
    fid = idx.upsert_file(item, "drive-1", kind, "indexed")
    idx.replace_chunks(fid, [dict(c, chunk_type="text") for c in chs])

st = idx.stats()
check("index populated", st["files"] == 3 and st["chunks"] == 4,
      f"{st['files']} files / {st['chunks']} chunks")

r = idx.search("possession schedule Cornell", k=5, mode="lexical")
check("BM25 finds the right doc",
      r["results"] and r["results"][0]["file"] == "land_review.pdf",
      r["results"][0]["file"] if r["results"] else "no hits")
check("citation carries page", r["results"][0]["page"] == 1)

r2 = idx.search("CloseDate", k=5, mode="lexical")
check("schema card is searchable",
      r2["results"] and r2["results"][0]["file"] == "closings_2026.csv")

r3 = idx.search('"retention class"', k=5, mode="lexical")
check("quoted phrase works", r3["results"] and
      r3["results"][0]["file"] == "standard.docx")

r4 = idx.search("owner; retention (class) *", k=5, mode="lexical")
check("punctuation does not break FTS", bool(r4["results"]))

r5 = idx.search("possession", k=5, mode="lexical", ext="pdf")
check("ext filter applied", all(x["ext"] == "pdf" for x in r5["results"]))

# incremental behaviour
fid, etag, status = idx.existing_etag("drive-1", "item-0")
check("etag round-trips", etag == "tag-0" and status == "indexed")

# fake vectors to exercise the dense path + RRF
import numpy as np


class FakeEmbedder:
    name, dim = "fake", 8

    def encode(self, texts, kind="passage"):
        v = np.zeros((len(texts), 8), dtype=np.float32)
        for i, t in enumerate(texts):
            for j, w in enumerate(["possession", "servicing", "retention", "owner",
                                   "closedate", "community", "cornell", "dataset"]):
                v[i, j] = t.lower().count(w)
        n = np.linalg.norm(v, axis=1, keepdims=True)
        n[n == 0] = 1
        return v / n


emb = FakeEmbedder()
rows = idx.pending_vector_chunks()
idx.store_vectors([r["chunk_id"] for r in rows],
                  emb.encode([r["text"] for r in rows]))
idx.meta_set("embed_backend", "fake")
idx.meta_set("embed_dim", 8)
st2 = idx.stats()
check("vectors stored", st2["vectors"] == 4 and st2["vector_coverage_pct"] == 100.0,
      f"{st2['vectors']} vectors")

r6 = idx.search("retention", k=5, embedder=emb, mode="hybrid")
check("hybrid returns fused hits", bool(r6["results"]) and
      r6["retrieval"]["vectors_used"] is True)
check("match provenance reported",
      r6["results"][0]["matched_by"] in ("both", "keyword", "semantic"),
      r6["results"][0]["matched_by"])

r7 = idx.search("servicing", k=5, embedder=emb, mode="semantic")
check("semantic-only path works", bool(r7["results"]) and
      r7["retrieval"]["lexical_hits"] == 0)

# re-chunking a file must not orphan vectors or FTS rows
fid0 = idx.conn().execute(
    "SELECT file_id FROM files WHERE item_id='item-0'").fetchone()["file_id"]
idx.replace_chunks(fid0, [{"text": "Rewritten content about grading permits.",
                           "page": 1, "chunk_type": "text"}])
st3 = idx.stats()
check("re-index replaces chunks cleanly", st3["chunks"] == 3,
      f"{st3['chunks']} chunks after replace")
r8 = idx.search("possession schedule", k=5, mode="lexical")
check("stale FTS rows removed",
      all(x["file"] != "land_review.pdf" for x in r8["results"]),
      str([x["file"] for x in r8["results"]]))
r9 = idx.search("grading permits", k=5, mode="lexical")
check("new content searchable", bool(r9["results"]) and
      r9["results"][0]["file"] == "land_review.pdf")

# deletion pruning
pruned = idx.delete_missing("drive-1", "", {"item-0", "item-1"})
check("missing files pruned", pruned == 1 and idx.stats()["files"] == 2,
      f"pruned {pruned}")

idx.conn().close()
idx._conn = None
idx.DB_PATH.unlink(missing_ok=True)
for sfx in ("-wal", "-shm"):
    idx.DB_PATH.with_name(idx.DB_PATH.name + sfx).unlink(missing_ok=True)

# --------------------------------------------------------------------------- #
passed = sum(1 for _, ok, _ in results if ok)
print(f"\n{passed}/{len(results)} checks passed")
sys.exit(0 if passed == len(results) else 1)
